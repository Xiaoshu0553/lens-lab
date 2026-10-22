from pathlib import Path
import json
import os
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FINAL = ROOT / 'LENS_Lab_评测平台完整功能逻辑与设计说明书.docx'
FIGURES = ROOT / 'figures'


def cjk_font_path():
    requested = os.environ.get('CJK_FONT')
    candidates = [requested] if requested else [
        '/System/Library/Fonts/STHeiti Light.ttc',
        '/System/Library/Fonts/STHeiti Medium.ttc',
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc',
        'C:/Windows/Fonts/msyh.ttc',
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise FileNotFoundError('未找到中文绘图字体，请将 CJK_FONT 设为本机可用的中文字体文件路径。')


FONT = cjk_font_path()

def font(size):
    return ImageFont.truetype(FONT, size)

def centered(draw, box, text, size=30, fill='#202A33'):
    f = font(size)
    lines = text.split('\n')
    heights = [draw.textbbox((0,0), line, font=f)[3] for line in lines]
    total = sum(heights) + (len(lines)-1)*12
    y = box[1] + (box[3]-box[1]-total)/2 - 3
    for line,h in zip(lines,heights):
        w = draw.textlength(line,font=f)
        draw.text(((box[0]+box[2]-w)/2,y),line,font=f,fill=fill)
        y += h+12

def node(draw, box, text, size=30, fill='#EDF2F6'):
    draw.rounded_rectangle(box, radius=16, fill=fill, outline='#AEBBC6',width=2)
    centered(draw,box,text,size)

def arrow(draw, a, b, width=4):
    draw.line((a,b),fill='#617383',width=width)
    import math
    angle=math.atan2(b[1]-a[1],b[0]-a[0])
    p1=(b[0]-17*math.cos(angle-.42),b[1]-17*math.sin(angle-.42))
    p2=(b[0]-17*math.cos(angle+.42),b[1]-17*math.sin(angle+.42))
    draw.polygon((b,p1,p2),fill='#617383')

def diagrams():
    FIGURES.mkdir(exist_ok=True)
    im=Image.new('RGB',(1800,700),'white'); d=ImageDraw.Draw(im)
    labels=['准备与版本冻结\n统一问卷 校园任务','资格确认与前测\n18–25岁 专本科\n统一30题','层内平衡分配\n四案例情景\nS 普通或 A 澄清','真实模型对话\n公共资料 完整历史']
    for i,label in enumerate(labels):
        x=30+i*447
        node(d,(x,65,x+390,225),label,30)
        if i<3: arrow(d,(x+391,135),(x+439,135))
    arrow(d,(1566,225),(1566,315))
    bottom=[('独立验证与改进\n新参与者或独立任务',30),('真实样本线下分析\n结果 差距 成本',477),('自动核对与导出\n保留版本及数据来源',924),('决策与首次提交\n参加80分 办理20分',1371)]
    for label,x in bottom: node(d,(x,315,x+390,465),label,29)
    for x in [1371,924,477]: arrow(d,(x,375),(x-57,375))
    centered(d,(120,510,1680,640),'合成展示与研究者试跑单独保存\n模拟数据不进入真实参与者的效应分析',31)
    im.save(FIGURES/'workflow.png')

    im=Image.new('RGB',(1800,1160),'white');d=ImageDraw.Draw(im)
    node(d,(190,25,1610,175),'两组共同输入\n公共基础提示 公开任务及统一资料 完整对话历史\n不含未提供的私事实 其他案例或内部评分',32)
    d.line([(900,175),(900,215),(435,215)],fill='#617383',width=4)
    arrow(d,(435,215),(435,260))
    d.line([(900,215),(1365,215)],fill='#617383',width=4)
    arrow(d,(1365,215),(1365,260))
    node(d,(60,260,810,425),'S 普通问答组\n基础提示 answer-dialogue-v6\n无额外主动澄清要求',33)
    node(d,(990,260,1740,425),'A 主动澄清组\n同一基础提示 + 额外澄清提示\nclarify-dialogue-v4 每轮保留',33)
    d.line([(435,425),(435,475),(900,475)],fill='#617383',width=4)
    d.line([(1365,425),(1365,475),(900,475)],fill='#617383',width=4)
    arrow(d,(900,475),(900,520))
    node(d,(240,520,1560,700),'同一真实模型直接生成自然语言回复\n回答本轮问题 或主动索取必要个案事实\n用户补充后 带全部历史继续回答原问题',34)
    arrow(d,(900,700),(900,765))
    node(d,(220,765,1580,920),'按实际输出记录两组行为\n识别并剥离共同标记【澄清提问】\n保存回复 动作 提示版本 服务耗时和异常',33)
    centered(d,(75,960,1725,1125),'S组自然追问也计数 标记不代表要求S组主动提问\nA组每轮最多问一个关键问题 没有累计两次上限\npilot/live不调用独立澄清模块 demo仅作离线规则模拟',30)
    im.save(FIGURES/'clarification.png')

def set_font(obj,size=None,bold=None,east='Songti SC',color='000000'):
    obj.font.name='Arial'
    if size is not None: obj.font.size=Pt(size)
    if bold is not None: obj.font.bold=bold
    obj.font.color.rgb=RGBColor.from_string(color)
    if hasattr(obj,'_element'):
        if obj._element.tag==qn('w:style'):
            rpr=obj._element.get_or_add_rPr()
        else:
            rpr=obj._element.get_or_add_rPr()
        rf=rpr.find(qn('w:rFonts'))
        if rf is None: rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
        for attr in list(rf.attrib):
            if 'theme' in attr.lower(): del rf.attrib[attr]
        rf.set(qn('w:eastAsia'),east)

def bookmark(p,name,idx):
    st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),str(idx));st.set(qn('w:name'),name)
    en=OxmlElement('w:bookmarkEnd');en.set(qn('w:id'),str(idx))
    p._p.insert(0,st);p._p.append(en)

def table(doc, rows):
    t=doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if len(rows[0])==4:
        widths=[2.3,4.4,5.1,4.5]
    elif rows[0][0].startswith('接口'):
        widths=[3.0,7.0,6.3]
    elif rows[0][0].startswith('字段'):
        widths=[5.0,5.5,5.8]
    elif rows[0][0]=='功能编号':
        widths=[1.8,6.7,7.8]
    else:
        widths=[3.2,6.6,6.5]
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for a,v in [('val','single'),('sz','5'),('color','D9D9D9')]: e.set(qn('w:'+a),v)
        borders.append(e)
    pr.append(borders)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','85'),('bottom','85'),('left','110'),('right','110')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5ECF2' if i==0 else ('F8FAFC' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.13
            p.paragraph_format.keep_with_next=False
            if j==0 and rows[0][0] in ['功能编号','角色','状态']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(txt);set_font(run,9.2,i==0,east='Heiti SC' if i==0 else 'Songti SC')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(4)
    set_font(p.add_run(''),4)

def add_toc_line(doc,title,number,bookmark_name):
    p=doc.add_paragraph(style='TOC Line')
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('w:anchor'),bookmark_name)
    r=OxmlElement('w:r');rpr=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts');rf.set(qn('w:eastAsia'),'Songti SC');rf.set(qn('w:ascii'),'Arial');rpr.append(rf)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'22');rpr.append(size)
    c=OxmlElement('w:color');c.set(qn('w:val'),'000000');rpr.append(c);r.append(rpr)
    text=OxmlElement('w:t');text.text=title;r.append(text);hyp.append(r);p._p.append(hyp)
    set_font(p.add_run('\t'+str(number) if number else '\t'),11)

def build():
    diagrams()
    source=(ROOT/'content.md').read_text()
    headings=[line[2:] for line in source.splitlines() if line.startswith('# ')]
    pmap=json.loads((ROOT/'page_map.json').read_text()) if (ROOT/'page_map.json').exists() else {}
    doc=Document();sec=doc.sections[0]
    for st in doc.styles:
        for border in st._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for name in ['Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name]
        st.font.italic=False
        ppr=st._element.get_or_add_pPr()
        for child in list(ppr):
            if child.tag==qn('w:numPr'):ppr.remove(child)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.25)
    sec.footer_distance=Cm(.85)
    normal=doc.styles['Normal'];set_font(normal,10.5)
    pf=normal.paragraph_format;pf.line_spacing=1.25;pf.space_after=Pt(5.5)
    pf.widow_control=True
    for name,size in [('Title',23),('Subtitle',16),('Heading 1',15.5),('Heading 2',12.2)]:
        st=doc.styles[name];set_font(st,size,True,east='Heiti SC');st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(15 if name=='Heading 1' else 10)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.15
    for name in ['List Paragraph']:
        set_font(doc.styles[name],10.5);doc.styles[name].paragraph_format.line_spacing=1.25
        doc.styles[name].paragraph_format.space_after=Pt(5)
    toc=doc.styles.add_style('TOC Line',1);set_font(toc,11)
    toc.paragraph_format.space_after=Pt(6);toc.paragraph_format.line_spacing=1.05
    sec.different_first_page_header_footer=True
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=foot.add_run();set_font(run,9,color='606060')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    doc.core_properties.title='LENS Lab 信息公平性评测平台完整功能逻辑与设计说明书'
    doc.core_properties.subject='挑战杯研究平台 功能 逻辑 目的 设计与验收'
    doc.core_properties.author='项目研究团队'
    doc.core_properties.keywords='信息公平性 AI信息素养 主动澄清 大学生 评测平台 挑战杯'

    p=doc.add_paragraph('LENS Lab 信息公平性评测平台',style='Title')
    p.paragraph_format.space_before=Pt(20)
    doc.add_paragraph('完整功能逻辑与设计说明书',style='Subtitle')
    p=doc.add_paragraph('面向语言表达差异的生成式人工智能信息公平性评测与自适应干预研究')
    set_font(p.runs[0],12,True,east='Heiti SC')
    p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(20)
    doc.add_paragraph('版本 university-v1    日期 2026 年 10 月 22 日')
    doc.add_paragraph('适用对象  项目负责人 研究成员 开发人员 指导教师与评审沟通人员')
    doc.add_paragraph('研究对象  18–25岁当前在读专科生和本科生')
    doc.add_paragraph('本文定义平台为什么建设、服务什么研究问题、包含哪些功能、如何执行实验、怎样形成可复核数据，以及正式投放前需要完成的验证。设计以当前信息素养前测与两组被试间实验为主线，同时明确原题目中语言表达差异的补充证据路径。')
    doc.add_paragraph('平台的主要价值是将研究方案落实为一致的实验流程与完整的原始记录。核心评价同时考察澄清必要性、任务有效性、信息公平性和交互成本；是否改善公平由真实研究数据回答。')
    doc.add_paragraph('阅读指引',style='Heading 2')
    doc.add_paragraph('理解项目定位与研究逻辑，先读第一至八章。梳理页面功能与指标，阅读第九至十四章。开展开发、投放和验收，阅读第十五至十八章及数据字典。第十九章提供一次完整任务的运行示例。')
    doc.add_paragraph('文中区分当前已实现、正式研究前须核验和后续扩展。功能说明以大学生版本实现为依据；测量有效性、任务可比性与研究结论须单独验证。126条合成展示记录不构成真实大学生调查样本。')
    doc.add_page_break()
    doc.add_paragraph('目录',style='Heading 1')
    for i,title in enumerate(headings,1):add_toc_line(doc,title,pmap.get(title,''),f'chapter_{i}')
    doc.add_page_break()

    lines=source.splitlines();i=0;idx=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('# '):
            idx+=1;p=doc.add_paragraph(line[2:],style='Heading 1');bookmark(p,f'chapter_{idx}',idx)
            if line.startswith(('# 附录二', '# 附录三')):p.paragraph_format.page_break_before=True
        elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 2')
        elif line.startswith('[[FIGURE:'):
            key=line[len('[[FIGURE:'):-2]
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            run=p.add_run();run.add_picture(str(FIGURES/(key+'.png')),width=Cm(16.3))
            pic=run._r.xpath('.//wp:docPr')[0];pic.set('descr','实验闭环流程' if key=='workflow' else '主动澄清逻辑分支')
            p=doc.add_paragraph('图 1  平台实验与研究闭环' if key=='workflow' else '图 2  真实模型的两组提示与记录流程')
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:set_font(r,9,color='505050')
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]);i+=1
            table(doc,rows);continue
        elif re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(line,style='List Paragraph')
            p.paragraph_format.left_indent=Cm(.45);p.paragraph_format.first_line_indent=Cm(-.45)
        else:doc.add_paragraph(line)
        i+=1
    for el in doc.element.body.xpath('.//w:pBdr'):
        el.getparent().remove(el)
    doc.save(FINAL)
    (ROOT/'headings.json').write_text(json.dumps(headings,ensure_ascii=False,indent=2))
    print(FINAL)
    print('source_characters',len(source),'chapters',len(headings),'paragraphs',len(doc.paragraphs),'tables',len(doc.tables))

if __name__=='__main__':build()
